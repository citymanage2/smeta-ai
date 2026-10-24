"""Фаза 8в: пометка «цена из КП» в скачиваемых xlsx оптимизации.

Строка с `kp_ref` заливается цветом «цена из КП», цена работы, которая
включает материалы, — отдельным цветом. Заливка — только оформление: значения
ячеек, формулы и итоги файла те же, что у тех же строк без `kp_ref`.
"""
import copy
import io

import openpyxl

from app.utils.kp_marks import KP_FILL_COLOR, KP_MATERIALS_FILL_COLOR, kp_fill, kp_ref_active

KP_REF = {
    "request_id": "r1", "offer_id": "o1", "supplier": "ООО Кабель",
    "offer_date": "2026-09-10", "offers_count": 2, "includes_materials": False,
    "field": "price_work", "price": 100.5,
}
KP_REF_MAT = {**KP_REF, "includes_materials": True, "price": 200}


def _rows():
    return [
        {"id": "s1", "type": "section", "name": "Раздел 1"},
        {"id": "a", "num": 1, "type": "work", "name": "Монтаж кабеля", "unit": "м",
         "qty": 10, "price_work": 100.5, "price_material": None, "kp_ref": KP_REF},
        {"id": "b", "num": 2, "type": "material", "name": "Кабель", "unit": "м",
         "qty": 10, "price_work": None, "price_material": 33.4665},
        {"id": "c", "num": 3, "type": "work", "name": "Монтаж лотка", "unit": "м",
         "qty": 3, "price_work": 200, "price_material": 50, "kp_ref": KP_REF_MAT},
    ]


def _strip(rows):
    out = copy.deepcopy(rows)
    for r in out:
        r.pop("kp_ref", None)
    return out


def _values(ws):
    return [[c.value for c in row] for row in ws.iter_rows()]


def _color(cell):
    fill = cell.fill
    if fill is None or fill.fill_type != "solid":
        return None
    return str(fill.start_color.rgb)[-6:]


def _export(rows, coefficient=None):
    from app.services.excel_service import generate_estimate_export
    data = generate_estimate_export(
        rows=rows, overhead_pct=10, transport_pct=5, contingency_pct=0,
        version_display_name="V1", coefficient=coefficient,
    )
    return openpyxl.load_workbook(io.BytesIO(data))


# ── kp_fill ─────────────────────────────────────────────────────────────────

def test_kp_fill_rules():
    assert kp_fill(None) is None
    assert kp_fill({}) is None
    assert kp_fill("строка") is None
    # Пометка без `field`/`price` (записана до правила) — недействующая.
    assert kp_fill({"request_id": "r1"}) is None
    assert kp_fill(KP_REF).start_color.rgb.endswith(KP_FILL_COLOR)
    assert kp_fill(KP_REF_MAT).start_color.rgb.endswith(KP_MATERIALS_FILL_COLOR)
    assert KP_FILL_COLOR != KP_MATERIALS_FILL_COLOR


def test_kp_ref_active_exact_numbers():
    row = {"type": "work", "price_work": 100.5, "price_material": 30, "kp_ref": KP_REF}
    assert kp_ref_active(row) is KP_REF
    # Разница меньше полкопейки — та же цена (float-шум после пересчёта).
    assert kp_ref_active({**row, "price_work": 100.504}) is KP_REF
    assert kp_ref_active({**row, "price_work": "100,5"}) is KP_REF
    # Цену поменяли — пометка больше не про эту цену.
    assert kp_ref_active({**row, "price_work": 100.51}) is None
    assert kp_ref_active({**row, "price_work": 120}) is None
    assert kp_ref_active({**row, "price_work": None}) is None
    # Смена цены другого поля пометку не снимает.
    assert kp_ref_active({**row, "price_material": 999}) is KP_REF
    # Без field/price, чужое поле, не объект — недействующая.
    legacy = {k: v for k, v in KP_REF.items() if k not in ("field", "price")}
    assert kp_ref_active({**row, "kp_ref": legacy}) is None
    assert kp_ref_active({**row, "kp_ref": {**KP_REF, "field": "cost"}}) is None
    assert kp_ref_active({**row, "kp_ref": {**KP_REF, "price": "100.5"}}) is None
    assert kp_ref_active({**row, "kp_ref": {**KP_REF, "price": True}}) is None
    assert kp_ref_active({**row, "kp_ref": "строка"}) is None
    assert kp_ref_active(None) is None


# ── generate_estimate_export ────────────────────────────────────────────────

def _row_by_name(ws, name):
    for row in ws.iter_rows(min_row=2):
        if row[2].value == name:
            return row
    raise AssertionError(name)


def test_estimate_export_fills_kp_rows():
    ws = _export(_rows()).active
    assert {_color(c) for c in _row_by_name(ws, "Монтаж кабеля")} == {KP_FILL_COLOR}
    assert {_color(c) for c in _row_by_name(ws, "Монтаж лотка")} == {KP_MATERIALS_FILL_COLOR}
    assert {_color(c) for c in _row_by_name(ws, "Кабель")} == {None}
    # Раздел сохраняет свою заливку.
    assert _color(_row_by_name(ws, "Раздел 1")[0]) == "BDD7EE"


def test_estimate_export_price_changed_after_kp_no_fill():
    rows = _rows()
    rows[1]["price_work"] = 120  # человек поменял цену строки после «Цены из КП»
    rows[3]["kp_ref"] = {k: v for k, v in KP_REF_MAT.items() if k not in ("field", "price")}
    ws = _export(rows).active
    assert {_color(c) for c in _row_by_name(ws, "Монтаж кабеля")} == {None}
    assert {_color(c) for c in _row_by_name(ws, "Монтаж лотка")} == {None}


def test_estimate_export_fill_checks_base_price_not_coefficient_price():
    """В `kp_ref.price` — исходная цена; коэффициент версии пометку не снимает."""
    ws = _export(_rows(), coefficient={"work": 1.1, "material": 1, "scope": "all"}).active
    row = _row_by_name(ws, "Монтаж кабеля")
    assert row[5].value == 110.55
    assert {_color(c) for c in row} == {KP_FILL_COLOR}


def test_estimate_export_totals_unchanged_by_kp_ref():
    with_kp = _export(_rows(), coefficient={"work": 1.1, "material": 1, "scope": "all"}).active
    without = _export(_strip(_rows()), coefficient={"work": 1.1, "material": 1, "scope": "all"}).active
    assert _values(with_kp) == _values(without)


def test_estimate_export_totals_exact_numbers():
    ws = _export(_rows()).active
    totals = {row[2].value: row[7].value for row in ws.iter_rows(min_row=2)
              if row[2].value and row[7].value is not None}
    # Работы: 10*100.5 + 3*200 = 1005 + 600 = 1605
    # Материалы: round(10*33.4665, 2) + 3*50 = 334.67 + 150 = 484.67
    assert totals["Работы (итого)"] == 1605.0
    assert round(totals["Материалы (итого)"], 2) == 484.67
    assert round(totals["Итого (базис)"], 2) == 2089.67
    assert round(totals["Накладные расходы (10%)"], 2) == 160.5
    assert round(totals["Транспортные расходы (5%)"], 2) == 24.23


# ── generate_statement_xlsx ─────────────────────────────────────────────────

COLUMNS = [
    {"key": "name", "label": "Наименование", "numeric": False},
    {"key": "qty", "label": "Кол-во", "numeric": True},
    {"key": "price_work", "label": "Цена работы", "numeric": True},
    {"key": "cost", "label": "Стоимость", "numeric": True},
]


def _statement_rows():
    return [
        {"_id": "a", "_kind": "work", "name": "Монтаж кабеля", "qty": 10,
         "price_work": 100.5, "cost": 1005, "__kp_ref": KP_REF},
        {"_id": "b", "_kind": "material", "name": "Кабель", "qty": 10,
         "price_work": None, "cost": 334.67},
        {"_id": "c", "_kind": "work", "name": "Монтаж лотка", "qty": 3,
         "price_work": 200, "cost": 750, "__kp_ref": KP_REF_MAT},
    ]


def _statement(rows):
    from app.utils.xlsx_statement import generate_statement_xlsx
    data = generate_statement_xlsx(COLUMNS, rows, title="", show_date=False)
    return openpyxl.load_workbook(io.BytesIO(data)).active


def _statement_row(ws, name):
    for row in ws.iter_rows():
        if row[0].value == name:
            return row
    raise AssertionError(name)


def test_statement_fills_kp_rows():
    ws = _statement(_statement_rows())
    assert {_color(c) for c in _statement_row(ws, "Монтаж кабеля")} == {KP_FILL_COLOR}
    assert {_color(c) for c in _statement_row(ws, "Монтаж лотка")} == {KP_MATERIALS_FILL_COLOR}
    assert {_color(c) for c in _statement_row(ws, "Кабель")} == {None}
    # Служебный ключ колонкой не становится.
    assert [c.value for c in _statement_row(ws, "Наименование")] == [
        "Наименование", "Кол-во", "Цена работы", "Стоимость",
    ]


def test_statement_totals_unchanged_by_kp_ref():
    rows = _statement_rows()
    stripped = copy.deepcopy(rows)
    for r in stripped:
        r.pop("__kp_ref", None)
    with_kp = _statement(rows)
    assert _values(with_kp) == _values(_statement(stripped))
    total = _statement_row(with_kp, "ИТОГО")
    assert total[3].value == 2089.67


def test_statement_kp_ref_in_regular_column_ignored():
    """Колонка заказчика с ключом «kp_ref» — не пометка: смотрим только служебный ключ."""
    rows = [{"_id": "a", "name": "Позиция", "qty": 1, "price_work": 1, "cost": 1,
             "kp_ref": KP_REF}]
    ws = _statement(rows)
    assert {_color(c) for c in _statement_row(ws, "Позиция")} == {None}
