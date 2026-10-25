"""Строка материала в скачиваемых xlsx — синий курсив, текст прижат вправо.

Правило жило только в перечне: в смете, версии оптимизации, сравнении версий и
разделах сводной материал выглядел так же, как работа. В сводной не узнавался и
тип строки — генератор сравнивал его с «Работа», а раздел хранит строки
документа с типом "work".

Оформление проверяется одним правилом на все генераторы: наименование
материала — синим курсивом и вправо, работы — нет. Кегль, гарнитура, рамки и
форматы чисел — вёрстка, их тест не держит. Оформление не трогает значения:
числа в ячейках закреплены отдельными тестами ниже.
"""
import io
from types import SimpleNamespace

import openpyxl
import pytest

from app.utils.estimate_rows import canonical_type
from app.utils.xlsx_row_style import MATERIAL_FONT_COLOR


def _book(data: bytes):
    return openpyxl.load_workbook(io.BytesIO(data))


def _blue_italic(cell) -> bool:
    color = cell.font.color.rgb if cell.font.color is not None else None
    return bool(cell.font.i) and isinstance(color, str) and color.endswith(MATERIAL_FONT_COLOR)


def _row_of(ws, name: str, name_col: int) -> int:
    for row in range(1, ws.max_row + 1):
        if ws.cell(row=row, column=name_col).value == name:
            return row
    raise AssertionError(f"строка «{name}» не найдена")


def _values(ws, row: int, columns) -> list:
    return [ws.cell(row=row, column=c).value for c in columns]


# ── Тип строки в любом написании ────────────────────────────────────────────

@pytest.mark.parametrize("raw, expected", [
    ("work", "work"), ("Работа", "work"), (" работы ", "work"),
    ("material", "material"), ("Материал", "material"), ("МАТЕРИАЛЫ", "material"),
    ("section", "section"), ("Раздел", "section"),
    (None, ""), ("", ""), ("Оборудование", "оборудование"),
])
def test_canonical_type(raw, expected):
    assert canonical_type(raw) == expected


# ── Данные для генераторов ──────────────────────────────────────────────────

# Раздел сводной — копия строк документа: тип "work"/"material".
DOC_ROWS = [
    {"type": "work", "name": "Вырубка деревьев", "unit": "шт",
     "qty": 303, "price_work": 3800, "price_material": None},
    {"type": "material", "name": "Рабочий", "unit": "чел/час",
     "qty": 606, "price_work": None, "price_material": 320},
]
# Старый снимок раздела — позиции ИИ: «Работа»/«Материал».
LEGACY_ROWS = [
    {"type": "Работа", "name": "Вырубка деревьев", "unit": "шт",
     "quantity": 303, "work_price": 3800, "material_price": None},
    {"type": "Материал", "name": "Рабочий", "unit": "чел/час",
     "quantity": 606, "work_price": None, "material_price": 320},
]

ITEMS = [
    {"type": "Работа", "name": "Монтаж кабеля", "unit": "м", "quantity": 10,
     "work_price": 100, "material_price": None, "price_list_name": "Прайс", "notes": "п1"},
    {"type": "Материал", "name": "Кабель ВВГнг 3х2,5", "unit": "м", "quantity": 10,
     "work_price": None, "material_price": 33.5, "price_list_name": "ИИ", "notes": "п2"},
    # Позиция ИИ в нижнем регистре — тоже материал.
    {"type": "материал", "name": "Хомут", "unit": "шт", "quantity": 4,
     "work_price": None, "material_price": 5, "price_list_name": "", "notes": ""},
]

VERSION_ROWS = [
    {"id": "s1", "type": "section", "name": "Раздел 1"},
    {"id": "a", "num": 1, "type": "work", "name": "Монтаж кабеля", "unit": "м",
     "qty": 10, "price_work": 100.5, "price_material": None,
     "kp_ref": {"field": "price_work", "price": 100.5, "includes_materials": False}},
    {"id": "b", "num": 2, "type": "material", "name": "Кабель", "unit": "м",
     "qty": 10, "price_work": None, "price_material": 33.5},
]

# Колонки и строки — как их собирает «Ведомость по сводной»
# (`SummaryEditorTabs.exportRows`): тип строки едет ключом `_kind`.
STATEMENT_COLUMNS = [
    {"key": "section_name", "label": "Раздел", "numeric": False},
    {"key": "num", "label": "№", "numeric": True},
    {"key": "name", "label": "Наименование", "numeric": False},
    {"key": "unit", "label": "Ед. изм.", "numeric": False},
    {"key": "qty", "label": "Кол-во", "numeric": True},
    {"key": "price_work", "label": "Цена работ", "numeric": True},
    {"key": "cost_work", "label": "Стоим. работ", "numeric": True},
    {"key": "price_material", "label": "Цена матер.", "numeric": True},
    {"key": "cost_material", "label": "Стоим. матер.", "numeric": True},
]
STATEMENT_ROWS = [
    {"_id": "c:1", "_kind": "work", "_section": "c", "section_name": "Гостиница",
     "num": 1, "name": "Вырубка деревьев", "unit": "шт", "qty": 303,
     "price_work": 3800, "cost_work": 1151400, "price_material": None, "cost_material": None},
    {"_id": "c:2", "_kind": "material", "_section": "c", "section_name": "Гостиница",
     "num": 2, "name": "Рабочий", "unit": "чел/час", "qty": 606,
     "price_work": None, "cost_work": None, "price_material": 320, "cost_material": 193920},
]


def _summary_sheet(rows):
    from app.utils.xlsx_summary import generate_summary_xlsx
    summary = SimpleNamespace(
        sections=[{"card_name": "Гостиница", "rows": rows}], overrides={},
    )
    return _book(generate_summary_xlsx(summary))["Гостиница"]


def _estimate_sheet(items):
    from app.utils.xlsx_exporter import ESTIMATE_SHEET_TITLE, generate_estimate_xlsx
    data, _ = generate_estimate_xlsx(items, overhead_pct=0, transport_pct=0)
    return _book(data)[ESTIMATE_SHEET_TITLE]


def _with_source_no(items):
    return [dict(item, source_no=str(n)) for n, item in enumerate(items, 1)]


def _version_sheet():
    from app.services.excel_service import generate_estimate_export
    return _book(generate_estimate_export(
        rows=VERSION_ROWS, overhead_pct=0, transport_pct=0, contingency_pct=0,
        version_display_name="V1",
    ))["Смета"]


def _comparison_sheet():
    from app.services.excel_service import generate_comparison_export
    versions = [
        {"id": vid, "version_display_name": vid, "rows": VERSION_ROWS,
         "overhead_pct": 0, "transport_pct": 0, "contingency_pct": 0}
        for vid in ("V0", "V1")
    ]
    return _book(generate_comparison_export(versions))["Сравнение"]


def _statement_sheet(columns=STATEMENT_COLUMNS, rows=STATEMENT_ROWS):
    from app.utils.xlsx_statement import generate_statement_xlsx
    return _book(generate_statement_xlsx(columns, rows, title="Ведомость по сводной")).worksheets[0]


def _label_statement_sheet():
    """Документ с колонками из файла заказчика: ключи свои, решают подписи."""
    columns = [
        {"key": "c0", "label": "№ п/п", "numeric": False},
        {"key": "c1", "label": "Наименование работ и затрат", "numeric": False},
        {"key": "c2", "label": "Примечание", "numeric": False},
    ]
    rows = [
        {"_id": "1", "_kind": "work", "c0": "1", "c1": "Монтаж кабеля", "c2": ""},
        {"_id": "2", "_kind": "material", "c0": "2", "c1": "Кабель", "c2": "из КП"},
    ]
    return _statement_sheet(columns, rows)


def _perechen_sheet():
    from app.services.excel_service import generate_list
    return _book(generate_list([
        {"type": "Работа", "name": "Монтаж кабеля", "unit": "м", "quantity": 10},
        {"type": "Материал", "name": "Кабель", "unit": "м", "quantity": 10},
    ])).worksheets[0]


# Генератор → (лист, колонка наименования, наименование работы, наименование материала).
GENERATORS = {
    "сводная, строки документа": lambda: (_summary_sheet(DOC_ROWS), 2, "Вырубка деревьев", "Рабочий"),
    "сводная, старый снимок": lambda: (_summary_sheet(LEGACY_ROWS), 2, "Вырубка деревьев", "Рабочий"),
    "смета": lambda: (_estimate_sheet(ITEMS), 2, "Монтаж кабеля", "Кабель ВВГнг 3х2,5"),
    "смета, тип ИИ строчными": lambda: (_estimate_sheet(ITEMS), 2, "Монтаж кабеля", "Хомут"),
    "смета с № исходной": lambda: (
        _estimate_sheet(_with_source_no(ITEMS)), 3, "Монтаж кабеля", "Кабель ВВГнг 3х2,5"),
    "версия оптимизации": lambda: (_version_sheet(), 3, "Монтаж кабеля", "Кабель"),
    "сравнение версий": lambda: (_comparison_sheet(), 3, "Монтаж кабеля", "Кабель"),
    "ведомость": lambda: (_statement_sheet(), 3, "Вырубка деревьев", "Рабочий"),
    "ведомость по подписям колонок": lambda: (_label_statement_sheet(), 2, "Монтаж кабеля", "Кабель"),
    "перечень": lambda: (_perechen_sheet(), 3, "Монтаж кабеля", "Кабель"),
}


@pytest.mark.parametrize("build", GENERATORS.values(), ids=GENERATORS.keys())
def test_material_row_is_blue_italic_and_right_work_is_not(build):
    ws, name_col, work_name, material_name = build()
    material = ws.cell(row=_row_of(ws, material_name, name_col), column=name_col)
    work = ws.cell(row=_row_of(ws, work_name, name_col), column=name_col)

    assert _blue_italic(material)
    assert material.alignment.horizontal == "right"
    assert not work.font.i
    assert work.alignment.horizontal != "right"


# ── Значения: оформление их не трогает ──────────────────────────────────────

@pytest.mark.parametrize("rows", [DOC_ROWS, LEGACY_ROWS], ids=["doc", "legacy"])
def test_summary_section_values(rows):
    """303 × 3800 = 1 151 400; 606 × 320 = 193 920 — в обоих форматах строк."""
    ws = _summary_sheet(rows)
    assert _values(ws, 2, range(1, 9)) == [
        1, "Вырубка деревьев", "шт", 303, 3800, 1151400, None, None,
    ]
    assert _values(ws, 3, range(1, 9)) == [
        2, "Рабочий", "чел/час", 606, None, None, 320, 193920,
    ]


@pytest.mark.parametrize("with_source_no", [False, True], ids=["plain", "source-no"])
def test_estimate_material_values(with_source_no):
    items = _with_source_no(ITEMS) if with_source_no else ITEMS
    ws = _estimate_sheet(items)
    off = 1 if with_source_no else 0
    cable = _row_of(ws, "Кабель ВВГнг 3х2,5", 2 + off)
    assert _values(ws, cable, range(2 + off, 11 + off)) == [
        "Кабель ВВГнг 3х2,5", "м", 10, None, None, 33.5, 335, "ИИ", "п2",
    ]


def test_version_export_material_values():
    ws = _version_sheet()
    assert _values(ws, _row_of(ws, "Кабель", 3), range(1, 9)) == [
        2, "Материал", "Кабель", "м", 10, None, 33.5, 335,
    ]


def test_statement_number_is_integer():
    """№ — целым числом, а не «1,00»."""
    ws = _statement_sheet()
    number = ws.cell(row=_row_of(ws, "Вырубка деревьев", 3), column=2).value
    assert number == 1
    assert isinstance(number, int)


def test_statement_material_values():
    ws = _statement_sheet()
    assert _values(ws, _row_of(ws, "Рабочий", 3), range(1, 10)) == [
        "Гостиница", 2, "Рабочий", "чел/час", 606, None, None, 320, 193920,
    ]


def test_statement_total_sums_costs_not_numbers_or_prices():
    """Итог ведомости: стоимости складываются, № и цены — нет."""
    ws = _statement_sheet()
    total = _row_of(ws, "ИТОГО", 1)
    # № (2), цена работ (6), стоимость работ (7), цена матер. (8), стоимость матер. (9).
    assert _values(ws, total, (2, 6, 7, 8, 9)) == [None, None, 1151400, None, 193920]
