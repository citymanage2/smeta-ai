"""Цена работы из КП в расчёте сметы и защита зарезервированного имени «КП: …».

План: plans/2026-09-14-zapros-kp.md, фаза 7 (критерий 26, проблемы 22, 23, 27).

Решение пользователя 14.09.2026 поверх решения 03.08.2026: у работы в расчёт
идёт минимум по ценам «КП: <поставщик>»; нет цен из КП — как раньше: «Из смет»,
затем минимум по подрядчикам. Раз имя «КП: …» даёт приоритет, записать его
может только запрос КП: загрузка прайса файлом и ручная правка каталога такое
имя подрядчика не принимают.
"""
import io
from typing import Optional

import openpyxl
import pytest
import pytest_asyncio
from sqlalchemy import delete, select

from app.models.price import PriceMaterial, PriceWork
from app.models.price_cache import PriceCacheWork
from app.services import price_service
from app.utils.auth import create_access_token
from app.utils.price_min import (
    ESTIMATE_CONTRACTOR,
    KP_CONTRACTOR_PREFIX,
    compute_min_price,
    is_reserved_contractor,
    kp_contractor,
)

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# ---------------------------------------------------------------------------
# Правило выбора цены работы
# ---------------------------------------------------------------------------

def test_prefix_is_fixed():
    assert KP_CONTRACTOR_PREFIX == "КП: "
    assert kp_contractor("  ООО Стройка ") == "КП: ООО Стройка"


def test_kp_price_beats_estimate_and_contractors():
    assert compute_min_price({ESTIMATE_CONTRACTOR: 900, "КП: А": 1000, "Подрядчик": 800}) == 1000


def test_minimum_among_kp_prices():
    assert compute_min_price({"КП: А": 1000, "КП: Б": 950}) == 950


def test_without_kp_old_rule_stays():
    assert compute_min_price({ESTIMATE_CONTRACTOR: 900, "Подрядчик": 800}) == 900
    assert compute_min_price({"Подрядчик": 800, "Другой": 700}) == 700


def test_empty_kp_price_does_not_count():
    """Ноль и мусор у «КП: …» ценой не считаются — правило идёт дальше."""
    assert compute_min_price({"КП: А": 0, "КП: Б": None, ESTIMATE_CONTRACTOR: 900}) == 900


def test_reserved_names():
    assert is_reserved_contractor("КП: Х")
    assert is_reserved_contractor("  кп:Х")
    assert not is_reserved_contractor("КПД Строй")
    assert not is_reserved_contractor(ESTIMATE_CONTRACTOR)


# ---------------------------------------------------------------------------
# Пути записи прайса не принимают «КП: …»
# ---------------------------------------------------------------------------

def _admin() -> dict:
    return {"Authorization": f"Bearer {create_access_token(1, 'admin')}"}


@pytest_asyncio.fixture(autouse=True)
async def _clean(db_session, monkeypatch):
    async def _none(_name: str) -> Optional[list]:
        return None

    def _batch(names, input_type="search_document"):
        return [None] * len(names)

    monkeypatch.setattr("app.routers.prices_catalog._generate_embedding_safe", _none)
    monkeypatch.setattr("app.routers.admin.generate_embeddings_batch", _batch)

    async def _wipe():
        await db_session.execute(delete(PriceWork))
        await db_session.execute(delete(PriceMaterial))
        await db_session.execute(delete(PriceCacheWork))
        await db_session.commit()

    await _wipe()
    yield
    await _wipe()


def _xlsx(sheets: dict) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for row in rows:
            ws.append(list(row))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.mark.asyncio
async def test_works_upload_with_kp_contractor_is_422(async_client, db_session):
    data = _xlsx({"Работы": [["Наименование", "Ед. изм.", "КП: Х"], ["Кладка", "м2", 500]]})
    r = await async_client.post(
        "/admin/price-lists/works", files={"file": ("p.xlsx", data, XLSX)}, headers=_admin(),
    )
    assert r.status_code == 422, r.text
    assert "КП: Х" in r.json()["detail"]
    db_session.expire_all()
    assert (await db_session.execute(select(PriceWork))).scalars().all() == []


@pytest.mark.asyncio
async def test_combined_upload_with_kp_contractor_is_422_and_keeps_price(async_client, db_session):
    db_session.add(PriceWork(name="Демонтаж", unit="м2", prices={"А": 100.0}, min_price=100.0))
    await db_session.commit()
    data = _xlsx({
        "Работы": [["Наименование", "Ед. изм.", "кп: Y"], ["Кладка", "м2", 500]],
        "Материалы": [["Наименование", "Ед. изм.", "Цена"], ["Кирпич", "шт", 25]],
    })
    r = await async_client.post(
        "/admin/prices/upload", files={"file": ("all.xlsx", data, XLSX)}, headers=_admin(),
    )
    assert r.status_code == 422, r.text
    db_session.expire_all()
    names = [w.name for w in (await db_session.execute(select(PriceWork))).scalars().all()]
    assert names == ["Демонтаж"], "полная перезаливка не должна была начаться"


@pytest.mark.asyncio
async def test_catalog_create_with_kp_contractor_is_422(async_client):
    r = await async_client.post(
        "/prices/catalog/works", headers=_admin(),
        json={"name": "Кладка", "unit": "м2", "prices": {"КП: Х": 500}},
    )
    assert r.status_code == 422, r.text


@pytest.mark.asyncio
async def test_catalog_edit_keeps_existing_kp_price_but_rejects_new_or_changed(async_client, db_session):
    work = PriceWork(name="Кладка", unit="м2", prices={"КП: А": 1000.0, "Подрядчик": 800.0},
                     min_price=1000.0)
    db_session.add(work)
    await db_session.commit()
    await db_session.refresh(work)
    url = f"/prices/catalog/works/{work.id}"

    # Форма каталога отправляет весь набор цен — существующая цена из КП не мешает правке.
    r = await async_client.put(url, headers=_admin(),
                               json={"name": "Кладка стен", "prices": {"КП: А": 1000.0, "Подрядчик": 750.0}})
    assert r.status_code == 200, r.text
    assert r.json()["price"] == 1000.0

    r = await async_client.put(url, headers=_admin(), json={"prices": {"КП: А": 900.0, "Подрядчик": 750.0}})
    assert r.status_code == 422
    r = await async_client.put(url, headers=_admin(),
                               json={"prices": {"КП: А": 1000.0, "КП: Б": 500.0, "Подрядчик": 750.0}})
    assert r.status_code == 422

    # Убрать цену из КП вручную можно.
    r = await async_client.put(url, headers=_admin(), json={"prices": {"Подрядчик": 750.0}})
    assert r.status_code == 200
    assert r.json()["price"] == 750.0


# ---------------------------------------------------------------------------
# Эталон стирает и цены из КП — и показывает это до записи (проблема 27)
# ---------------------------------------------------------------------------

@pytest.mark.asyncio
async def test_reference_preview_lists_kp_prices_among_removed(async_client, admin_token, db_session, monkeypatch):
    async def _noop_cache(db):
        return None

    monkeypatch.setattr(price_service, "load_cache", _noop_cache)
    monkeypatch.setattr(price_service, "duplicate_vectors_ready", lambda kind: False)
    db_session.add(PriceWork(
        name="Кладка стен", unit="м3",
        prices={"КП: ООО Рога": 1200.0, ESTIMATE_CONTRACTOR: 900.0}, min_price=1200.0,
    ))
    await db_session.commit()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["№", "Тип", "Наименование", "Ед. изм.", "Кол-во",
               "Цена работы (за ед.)", "Цена материала (за ед.)"])
    ws.append([1, "Работа", "Кладка стен", "м3", 1, 1000, None])
    buf = io.BytesIO()
    wb.save(buf)
    r = await async_client.post(
        "/prices/reference/preview",
        files={"file": ("смета.xlsx", buf.getvalue(), XLSX)},
        headers={"Authorization": admin_token},
    )
    assert r.status_code == 200, r.text
    entry = r.json()["plan"][0]
    assert entry["action"] == "reprice"
    assert {"contractor": "КП: ООО Рога", "price": 1200.0} in entry["removed"]
