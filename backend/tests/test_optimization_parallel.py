"""Оптимизация сметы «найти цену ниже»: что попадает в файл и как она возобновляется.

Фоновая оптимизация (`_run_optimization_background`, job `task.optimize`) ищет
по каждой выбранной позиции цену ниже — прайс, затем веб-поиск — и дописывает
результат в файл: колонки «Цена сниженная» на листах сметы и лист «Сравнение»
(было / стало / экономия). Прогон настоящий, на тестовой БД: прайс — строки в
БД, файл — в S3, подменён только ИИ (веб-поиск). Результат читается из файла
слота `optimized` — его скачивает пользователь.

8.7: чекпоинт `partial_results` — повторный запуск не ищет уже найденное, ключ —
     пара (лист, строка).
8.8: сбой поиска по одной позиции не роняет оптимизацию.
"""
import io

import openpyxl

from app.routers.tasks import _run_optimization_background
from app.services import storage_service

PROMPT = ""


def _estimate_xlsx(*sheets: str) -> bytes:
    """Смета: на каждом листе шапка и позиции со строки 2."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title in sheets or ("Смета",):
        ws = wb.create_sheet(title)
        ws.append(["№", "Тип", "Наименование", "Ед. изм.", "Кол-во", "Цена работы (за ед.)"])
        ws.append([1, "Работа", "Кладка кирпичная", "м3", 10, 820])
        ws.append([2, "Материал", "Кирпич М150", "шт", 1000, 10])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _item(name: str, type_: str, row_index: int, price_incl_vat: float, sheet: str = "Смета") -> dict:
    """Позиция так, как её присылает экран оптимизации (`OptimizeItem`)."""
    return {
        "sheet": sheet, "row_index": row_index, "name": name, "type": type_,
        "quantity": 10, "unit": "м3", "price_excl_vat": price_incl_vat,
        "price_incl_vat": price_incl_vat, "total": price_incl_vat * 10,
    }


def _number(value):
    return value if isinstance(value, (int, float)) else None


async def _optimize(processor_env, app_sessions, items, estimate: bytes, **task_fields):
    """Запустить оптимизацию; вернуть задачу и лист «Сравнение» как строки
    (наименование, цена было, цена стало, экономия на ед., экономия %, источник)."""
    task = await processor_env.task(
        "ESTIMATE_FROM_LIST", status="processing", estimation_status="optimizing", **task_fields,
    )
    await processor_env.price_db.load()

    await _run_optimization_background(str(task.id), items, PROMPT, estimate, app_sessions)

    (result,) = await processor_env.results(task, slot="optimized")
    wb = openpyxl.load_workbook(io.BytesIO(await storage_service.load_bytes(result.storage_key)))
    task = await processor_env.reload(task)
    rows = [
        (name, before, _number(after), saving, pct, source)
        for name, before, after, saving, pct, source in wb["Сравнение"].iter_rows(
            min_row=2, values_only=True,
        )
        if name != "ИТОГО экономия"
    ]
    return task, rows, wb


async def test_lower_price_gives_savings_and_higher_price_is_not_offered(
    processor_env, app_sessions, price_db, claude_api,
):
    """Цена ниже — в файл с экономией; цена выше — не предлагается вовсе."""
    price_db.work("Кладка кирпичная", unit="м3", prices={"Подрядчик": 800})
    price_db.material("Кирпич М150", unit="шт", price=15)
    ai = claude_api()  # обе цены нашлись в прайсе — веб-поиска нет

    task, rows, wb = await _optimize(processor_env, app_sessions, [
        _item("Кладка кирпичная", "work", 2, 1000),
        _item("Кирпич М150", "material", 3, 12),
    ], _estimate_xlsx())

    assert task.status == "completed"
    assert task.estimation_status == "optimized"
    assert [row[:5] for row in rows] == [
        ("Кладка кирпичная", 1000, 800, 200, 20),
        ("Кирпич М150", 12, None, None, None),
    ]
    # Сниженная стоимость строки сметы: 800 × 10.
    sheet = wb["Смета"]
    header = [c.value for c in sheet[1]]
    assert sheet.cell(row=2, column=header.index("Стоимость сниженная") + 1).value == 8000
    assert ai.calls == []


async def test_search_error_on_one_item_does_not_stop_optimization(
    processor_env, app_sessions, price_db, claude_api,
):
    """Строка прайса без цены роняет поиск по своей позиции — остальные считаются,
    задача завершается, а позиция помечена ошибкой поиска."""
    price_db.work("Кладка кирпичная", unit="м3", prices={})  # цены нет — поиск падает
    price_db.material("Кирпич М150", unit="шт", price=9)
    claude_api()

    task, rows, _ = await _optimize(processor_env, app_sessions, [
        _item("Кладка кирпичная", "work", 2, 1000),
        _item("Кирпич М150", "material", 3, 12),
    ], _estimate_xlsx())

    assert task.status == "completed"
    assert rows == [
        ("Кладка кирпичная", 1000, None, None, None, "Ошибка поиска"),
        ("Кирпич М150", 12, 9, 3, 25, "Прайс-лист"),
    ]


async def test_resume_skips_done_items_by_sheet_and_row(
    processor_env, app_sessions, price_db, claude_api,
):
    """Строка 2 есть на обоих листах; посчитана только на первом — второй лист
    считается, первый берётся из чекпоинта и повторно не ищется."""
    price_db.work("Кладка кирпичная", unit="м3", prices={"Подрядчик": 800})
    claude_api()
    done = {
        "sheet": "Раздел 1", "row_index": 2, "name": "Кладка кирпичная",
        "original_price": 1000, "new_price": 777, "source": "Прайс-лист",
        "savings_abs": 223, "savings_pct": 22.3, "has_vat": True,
    }

    task, rows, _ = await _optimize(
        processor_env, app_sessions,
        [
            _item("Кладка кирпичная", "work", 2, 1000, sheet="Раздел 1"),
            _item("Кладка кирпичная", "work", 2, 1000, sheet="Раздел 2"),
        ],
        _estimate_xlsx("Раздел 1", "Раздел 2"),
        progress_data={"partial_results": [done]},
    )

    assert task.status == "completed"
    assert [row[2] for row in rows] == [777, 800]
    assert task.progress_data is None  # чекпоинт больше не нужен

