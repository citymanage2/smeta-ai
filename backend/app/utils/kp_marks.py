"""Пометка строки «цена из КП» в скачиваемых xlsx оптимизации.

Строка, цена которой пришла из КП (`kp_ref`, пишет `services/kp_apply.py`),
заливается цветом «цена из КП». Цена работы, которая включает материалы, —
отдельным цветом: человеку надо проверить материал позиции, иначе он оплачен
дважды. Это только оформление — значения ячеек и итоги файла заливка не меняет.

**Пометка живёт, пока цена та же.** `kp_apply` пишет в `kp_ref` поле цены
(`field`: `price_work` | `price_material`) и записанную исходную цену (`price`,
без коэффициента). Человек поменял цену строки — пометка «цена из КП» больше не
про эту цену и не действует (`kp_ref_active`). Пометка без `field`/`price` —
недействующая. Фронт проверяет зеркально по тем же ключам.

Правило одно на оба генератора (`excel_service.generate_estimate_export`,
`xlsx_statement.generate_statement_xlsx`), чтобы файлы не расходились в цвете.
"""
from __future__ import annotations

from typing import Any, Optional

from openpyxl.styles import PatternFill

from app.utils.price_coercion import coerce_price

# Голубой — цена из КП; персиковый — цена работы включает материалы.
KP_FILL_COLOR = "DDEBF7"
KP_MATERIALS_FILL_COLOR = "FCE4D6"

_KP_FILL = PatternFill(start_color=KP_FILL_COLOR, end_color=KP_FILL_COLOR, fill_type="solid")
_KP_MATERIALS_FILL = PatternFill(
    start_color=KP_MATERIALS_FILL_COLOR, end_color=KP_MATERIALS_FILL_COLOR, fill_type="solid",
)

# В ведомости ключи строк — это ключи колонок документа, и у заказчика может
# быть своя колонка «kp_ref». Поэтому пометка едет служебным ключом.
STATEMENT_KP_KEY = "__kp_ref"


KP_PRICE_FIELDS = ("price_work", "price_material")
# Полкопейки: цена та же, если разница — шум округления float.
KP_PRICE_EPS = 0.005


def _ref_price(kp_ref: Any) -> Optional[tuple[str, float]]:
    """(поле, цена) пометки; None — пометка не объект или без `field`/`price`."""
    if not isinstance(kp_ref, dict):
        return None
    field = kp_ref.get("field")
    price = kp_ref.get("price")
    if field not in KP_PRICE_FIELDS:
        return None
    if isinstance(price, bool) or not isinstance(price, (int, float)):
        return None
    return field, float(price)


def kp_ref_active(row: Any) -> Optional[dict]:
    """`kp_ref` строки, если он действует: цена в `row[field]` та же, что записана.

    Строка — в исходных ценах (без коэффициента), как её хранит версия.
    """
    if not isinstance(row, dict):
        return None
    kp_ref = row.get("kp_ref")
    ref = _ref_price(kp_ref)
    if ref is None:
        return None
    field, price = ref
    value = coerce_price(row.get(field))
    if value is None or abs(value - price) >= KP_PRICE_EPS:
        return None
    return kp_ref


def kp_fill(kp_ref: Any) -> Optional[PatternFill]:
    """Заливка по уже проверенной пометке (`kp_ref_active`); нет пометки,
    она не объект или без `field`/`price` — None."""
    if _ref_price(kp_ref) is None:
        return None
    if kp_ref.get("includes_materials") is True:
        return _KP_MATERIALS_FILL
    return _KP_FILL
